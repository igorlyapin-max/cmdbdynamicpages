"""Generate AA delivery workbooks and verify them without local skill dependencies."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import zipfile

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
FLOW = re.compile(r"\b(?:OAPI|AAPI|IF|L|M|H)(?:0|[1-9][0-9]*)\b")
SECRET = re.compile(
    r"-----BEGIN(?: [A-Z]+)? PRIVATE KEY-----|"
    r"\b(?:password|passwd|token|secret|api[_-]?key)\s*=\s*\S+|"
    r"\b(?:bearer|basic)\s+[A-Za-z0-9._~+/-]+=*|"
    r"\b(?:ghp_|glpat-|sk-|xoxb-)[A-Za-z0-9_-]{6,}", re.I,
)
BOOKS = {
    "healthcheck-map": ("Карта HealthCheck(Шаблон) (2).xlsx", "healthcheck-map.md", "healthchecks"),
    "metrics-map": ("Карта метрик(Шаблон).xlsx", "metrics-map.md", "metrics"),
    "event-logging-map": ("Карта регистрации событий.xlsx", "event-logging-map.md", "events"),
    "secrets-rotation-map": ("Смена секретов v1.1(шаблон).xlsx", "secrets-map.md", "secrets"),
}
HEADERS = {
    "Системы": ["Система/Cервис", "Информационный поток"],
    "HealthCheck": ["Система/Cервис", "Ресурс", "Тип ресурса", "Описание", "Статус"],
    "Метрики": ["Система/Cервис", "Метрика", "Label 1", "Label N", "Тип метрики", "Описание"],
    "События": ["Система/Cервис", "Тип события", "Дополнительная информация", "Требуется отражение результата", "ID корреляции"],
    "Смена секретов": ["Контур", "Информационный поток", "Имя секрета", "Применяется/\nне применяется", "Тип секрета", "Система/Cервис\nИнициатор\n(способ хранения секрета)", "Система/Cервис\nЦель\n(способ хранения секрета)", "Периодичность", "Способ обновления", "Ответственный"],
}
SHEETS = {
    "healthcheck-map": ["Системы", "HealthCheck", "Справочник типов ресурсов"],
    "metrics-map": ["Системы", "Метрики"],
    "event-logging-map": ["Системы", "События", "Стандарт разработки"],
    "secrets-rotation-map": ["Смена секретов"],
}
# Fingerprints of protected cells in the approved templates, not of generated data.
# Updating a template requires explicit canonical validation before changing these.
PROTECTED_CELLS = {
    "healthcheck-map": "f2d2e0f22551edad92c9b798653f87846553cec41a88cf60ef0d25684bbf86ab",
    "metrics-map": "a1f1d6bd0a83f5fd3fa855d19ab20992be5f6880e44646d2e547189234dc2048",
    "event-logging-map": "32749377dc74b2a804b455922bacd7f84839a95c2c6cdfb40bc3b5c25db17227",
    "secrets-rotation-map": "87ff0335359f916e4c0cde4c1dfd597e27374a88420197fb0236afefeaf18251",
}


def clean(value: str) -> str:
    return value.strip().replace("\\|", "|").replace("`", "").replace("<br>", "\n")


def table(root: Path, source: str, marker: str) -> list[dict[str, str]]:
    text = (root / "aa" / source).read_text(encoding="utf-8")
    token = f"<!-- aa-table: {marker} -->"
    if text.count(token) != 1:
        raise ValueError(f"{source}: expected one {token}")
    lines = text.split(token, 1)[1].strip().splitlines()
    rows = []
    for line in lines:
        if not line.startswith("|"):
            break
        rows.append([clean(v) for v in re.split(r"(?<!\\)\|", line.strip()[1:-1])])
    if len(rows) < 3 or not all(re.fullmatch(r":?-+:?", cell) for cell in rows[1]):
        raise ValueError(f"{source}: invalid/empty marked table {marker}")
    keys = rows[0]
    if len(set(keys)) != len(keys):
        raise ValueError(f"{source}: duplicate table columns")
    if any(len(row) != len(keys) for row in rows[2:]):
        raise ValueError(f"{source}: ragged table {marker}")
    return [dict(zip(keys, row, strict=True)) for row in rows[2:]]


def sheet(rows: list[list[str]], start: int = 4) -> dict:
    return {"start_row": start, "rows": [
        {chr(ord("B") + i): value for i, value in enumerate(row)} for row in rows
    ]}


def details(row: dict[str, str], excluded: set[str]) -> str:
    return "\n".join(f"{key}: {value}" for key, value in row.items() if key not in excluded)


def workbook_data(root: Path, name: str) -> dict:
    """Map authoritative Markdown rows to the canonical template columns."""
    _, source, marker = BOOKS[name]
    rows = table(root, source, marker)
    # Column names are explicit: schema drift must fail rather than drop information.
    if name == "healthcheck-map":
        systems = sorted({(r["Система"], r["Поток"]) for r in rows})
        data = {"Системы": sheet([list(r) for r in systems]), "HealthCheck": sheet([
            [r["Система"], r["Ресурс"], r["Тип ресурса"], details(r, {"Система", "Ресурс", "Тип ресурса", "Статус"}), r["Статус"]] for r in rows
        ])}
    elif name == "metrics-map":
        systems = sorted({(r["Система"], r["Поток"]) for r in rows})
        data = {"Системы": sheet([list(r) for r in systems]), "Метрики": sheet([
            [r["Система"], r["Метрика"], r["Labels"], "Нет дополнительных labels", r["Тип"], details(r, {"Система", "Метрика", "Labels", "Тип"})] for r in rows
        ])}
    elif name == "event-logging-map":
        systems = sorted({(r["Система"], r["Поток"]) for r in rows})
        data = {"Системы": sheet([list(r) + ["Требует согласования"] for r in systems]), "События": sheet([
            [r["Система"], r["Событие"], details(r, {"Система", "Событие", "Результат", "Корреляция"}), r["Результат"], r["Корреляция"]] for r in rows
        ], start=5)}
    else:
        keys = ["Контур", "Поток", "Секрет", "Применимость", "Тип", "Инициатор / хранение", "Получатель / хранение", "Периодичность", "Обновление", "Ответственный"]
        if set(rows[0]) != set(keys):
            raise ValueError("secrets table columns differ from the workbook contract")
        data = {"Смена секретов": sheet([[r[k] for k in keys] for r in rows])}
    return {"sheets": data}


def flow_ids(root: Path) -> set[str]:
    return set(re.findall(r"^\| ((?:OAPI|AAPI|IF|L|M|H)(?:0|[1-9][0-9]*)) \|", (root / "aa/information-model.md").read_text(encoding="utf-8"), re.M))


def validate_book(root: Path, name: str) -> None:
    path = root / "aa/xlsx" / f"{name}.xlsx"
    if not path.is_file():
        raise ValueError(f"missing workbook: {path}")
    expected = workbook_data(root, name)["sheets"]
    with zipfile.ZipFile(path) as archive:
        for entry in archive.infolist():
            if entry.file_size > 10 * 1024 * 1024:
                raise ValueError(f"{name}: oversized archive entry")
            if any(part in entry.filename for part in ("externalLinks", "connections", "vbaProject")):
                raise ValueError(f"{name}: prohibited external metadata")
            if entry.filename.endswith((".xml", ".rels")):
                body = archive.read(entry)
                if any(marker in body for marker in (b"absPath", b'TargetMode="External"', b"connections", b"vbaProject")):
                    raise ValueError(f"{name}: prohibited archive metadata")
    book = load_workbook(path, data_only=False)
    try:
        if book.sheetnames != SHEETS[name]:
            raise ValueError(f"{name}: sheet names differ")
        ids = flow_ids(root)
        for page in book:
            if page.max_row > 5000 or page.max_column > 50:
                raise ValueError(f"{name}: worksheet limits exceeded")
            for row in page:
                for cell in row:
                    if cell.data_type == "f" or cell.hyperlink:
                        raise ValueError(f"{name}: prohibited formula or hyperlink {cell.coordinate}")
                    if isinstance(cell.value, str):
                        if SECRET.search(cell.value):
                            raise ValueError(f"{name}: prohibited secret material")
                        if set(FLOW.findall(cell.value)) - ids:
                            raise ValueError(f"{name}: unknown flow at {page.title}!{cell.coordinate}")
        for title, contract in expected.items():
            page = book[title]
            headers = HEADERS[title] + (["Объект целевой системы"] if title == "Системы" and name == "event-logging-map" else [])
            for index, value in enumerate(headers, start=2):
                if page.cell(2, index).value != value:
                    raise ValueError(f"{name}: header changed in {title}")
            start = contract["start_row"]
            cells = {(f"{col}{start + i}"): value for i, row in enumerate(contract["rows"]) for col, value in row.items()}
            for address, value in cells.items():
                if page[address].value != value:
                    raise ValueError(f"{name}: {title}!{address} differs from Markdown")
            for row in page.iter_rows(min_row=start):
                for cell in row:
                    if cell.value is not None and cell.coordinate not in cells:
                        raise ValueError(f"{name}: extra cell {title}!{cell.coordinate}")
        protected = [
            (page.title, cell.coordinate, cell.value)
            for page in book for row in page for cell in row
            if cell.value is not None and (page.title not in expected or cell.row < expected[page.title]["start_row"])
        ]
        digest = hashlib.sha256(json.dumps(protected, ensure_ascii=False, separators=(",", ":"), default=str).encode()).hexdigest()
        if digest != PROTECTED_CELLS[name]:
            raise ValueError(f"{name}: canonical instructions/dictionary cells changed")
    finally:
        book.close()


def export(root: Path, output: Path) -> None:
    output.mkdir(parents=True, exist_ok=True)
    for name in BOOKS:
        (output / f"{name}.json").write_text(json.dumps(workbook_data(root, name), ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def generate(root: Path) -> None:
    templates = Path(os.environ.get("AA_XLSX_TEMPLATE_DIR", Path.home() / "projects/files/aa"))
    helper = Path(os.environ.get("AA_XLSX_HELPER", Path.home() / ".codex/skills/architecture-artifacts/scripts/aa_xlsx.py"))
    if not helper.is_file() or any(not (templates / b[0]).is_file() for b in BOOKS.values()):
        raise ValueError("canonical helper/templates unavailable: set AA_XLSX_HELPER and AA_XLSX_TEMPLATE_DIR")
    with tempfile.TemporaryDirectory(prefix="cmdp-aa-") as temporary:
        output = Path(temporary)
        export(root, output)
        # Validate all generated files before replacing the versioned set.
        for name, (template, _, _) in BOOKS.items():
            args = ["--template", str(templates / template), "--output", str(output / f"{name}.xlsx"), "--data", str(output / f"{name}.json"), "--information-model", str(root / "aa/information-model.md")]
            for command in ("create", "validate"):
                subprocess.run([sys.executable, str(helper), command, *args], check=True)
        (root / "aa/xlsx").mkdir(exist_ok=True)
        for name in BOOKS:
            (root / "aa/xlsx" / f"{name}.xlsx").write_bytes((output / f"{name}.xlsx").read_bytes())


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=["export", "generate", "validate"])
    parser.add_argument("--root", type=Path, default=ROOT)
    parser.add_argument("--output-dir", type=Path)
    args = parser.parse_args()
    try:
        if args.command == "export":
            if not args.output_dir:
                parser.error("export requires --output-dir")
            export(args.root, args.output_dir)
        elif args.command == "generate":
            generate(args.root)
        else:
            for name in BOOKS:
                validate_book(args.root, name)
                print(f"OK {name}: canonical layout and Markdown values")
    except (ValueError, KeyError, OSError, zipfile.BadZipFile, subprocess.CalledProcessError) as error:
        parser.exit(1, f"AA workbook error: {error}\n")


if __name__ == "__main__":
    main()
